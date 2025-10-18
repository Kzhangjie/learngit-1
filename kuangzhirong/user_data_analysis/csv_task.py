import pandas as pd
from dataclasses import dataclass
import json
import openpyxl
from openpyxl import Workbook

@dataclass
class task_info:
    session_id: str
    request_id: str
    user_id: str
    model: str
    input: str
    context: str
    request_body: str
    response_body: str
    reasoning_content: str
    tool_calls: str
    timestamp: str


if __name__ == "__main__":
    file_path = r"D:\lingxi_task\云文档问答2025-09-11T07-57_export.csv"
    df = pd.read_csv(file_path)
    task_info_list = []
    # 按session_id分组，每个session内部按model分组存储request_body列表
    session_dict = {}
    
    for index, row in df.iterrows():
        session_id = row["会话ID"]
        request_body = json.loads(row["请求体"])
        request_id = row["请求ID"]
        user_id = row["用户ID"]
        model = request_body.get("model", "")  # 从request_body中获取model
        
        # 如果session_id不存在，创建新的session字典
        if session_id not in session_dict:
            session_dict[session_id] = {}
        
        # 如果该session下的model不存在，创建新列表
        if model not in session_dict[session_id]:
            session_dict[session_id][model] = []
        
        # 将request_body添加到对应session和model的列表中
        session_dict[session_id][model].append(request_body)

    # 每个model列表只保留最后一个元素
    for session_id, models_dict in session_dict.items():
        for model in models_dict:
            if len(models_dict[model]) > 0:
                models_dict[model] = [models_dict[model][-1]]

    print(f"总共 {len(session_dict)} 个session")
    for session_id, models_dict in session_dict.items():
        print(f"Session: {session_id}")
        for model, request_bodies in models_dict.items():
            print(f"  Model: {model}, 数量: {len(request_bodies)}")
    
    # 保存为Excel文件
    wb = Workbook()
    wb.remove(wb.active)  # 删除默认的sheet
    
    for session_id, models_dict in session_dict.items():
        # 创建以session_id命名的sheet
        sheet_name = str(session_id)[:31]  # Excel sheet名称限制31个字符
        ws = wb.create_sheet(title=sheet_name)
        
        current_row = 1
        
        for model, request_bodies in models_dict.items():
            # 写入model名称
            ws.cell(row=current_row, column=1, value="Model")
            ws.cell(row=current_row, column=2, value=model)
            current_row += 1
            
            # 写入messages标题
            ws.cell(row=current_row, column=1, value="Messages")
            current_row += 1
            
            # 遍历request_body中的messages
            for request_body in request_bodies:
                messages = request_body.get("messages", [])
                for i, message in enumerate(messages):
                    # 将message以JSON格式保存到单元格
                    ws.cell(row=current_row, column=1, value=f"Message {i+1}")
                    ws.cell(row=current_row, column=2, value=json.dumps(message, ensure_ascii=False, indent=2))
                    current_row += 1
            
            current_row += 2  # model之间空两行
    
    # 保存Excel文件
    output_path = r"D:\lingxi_task\云文档问答_session_data_3.xlsx"
    wb.save(output_path)
    print(f"数据已保存到: {output_path}")
    
    print(task_info_list)