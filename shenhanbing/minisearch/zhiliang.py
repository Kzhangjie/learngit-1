# 生成回答的质量分析--chat/websearch各500条
#

import openpyxl
from utils.gateway_api_v2_stream import GateWayAPI
import sqlite3
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib


DB_PATH = "dialog_data.db"
MINI_TABLE = "Mini_all_Data"
OUTPUT_EXCEL = "text.xlsx"
sheetname = "Sheet1"
输出列名 = "判定"
输出列编号 = "D"

Prompt = """
# 任务
把要分析的内容，严格按分析标准进行分析，并必须按照输出格式进行输出，不能进行额外备注。

# 要分析的内容
问题：{{.question}}
回答：{{.answer}}

# 分析标准
结合问题，判定一下回答内容质量采用评分制度，满分3分，从0分开始计算
没有答非所问+1分，没有缺少意图+1分，没有幻觉（回答内容可信任）+1分

# 输出格式
请严格按照以下格式范例输出，以便程序自动提取结果，不能有任何备注与解释，只返回分数：
3或2或1或0
"""

model = "qwen-max"
version = ""
llm_arguments = {
    "temperature": 0.1,
    "max_tokens": 8000,
    "top_p": 0.1,
}
context = ""
sec_text = {}

CHAT_LIMIT = 500
WEBSEARCH_LIMIT = 500


def get_connection():
    return sqlite3.connect(DB_PATH)


def export_structured_records():
    conn = get_connection()
    df = pd.read_sql_query(f"SELECT * FROM {MINI_TABLE}", conn)
    conn.close()

    sort_cols = ["SessionID"]
    if "Timestamp" in df.columns:
        sort_cols.append("Timestamp")
    df = df.sort_values(by=sort_cols).reset_index(drop=True)

    chat_results = []
    websearch_results = []

    for session_id, group in df.groupby("SessionID"):
        group = group.reset_index(drop=True)
        i = 0
        while i + 2 < len(group) and len(chat_results) < CHAT_LIMIT:
            if (
                group.loc[i, "Role"] == "user"
                and group.loc[i, "Type"] == "text"
                and group.loc[i + 1, "Role"] == "assistant"
                and group.loc[i + 1, "Type"] == "code"
                and group.loc[i + 2, "Role"] == "assistant"
                and group.loc[i + 2, "Type"] == "text"
            ):
                chat_results.append(
                    {
                        "问题": group.loc[i, "Content"],
                        "回答": group.loc[i + 2, "Content"],
                        "分发": "chat",
                        "判定": "",
                    }
                )
                i += 3
            else:
                i += 1

    for session_id, group in df.groupby("SessionID"):
        group = group.reset_index(drop=True)
        i = 0
        while i + 4 < len(group) and len(websearch_results) < WEBSEARCH_LIMIT:
            if (
                group.loc[i, "Role"] == "user"
                and group.loc[i, "Type"] == "text"
                and group.loc[i + 1, "Role"] == "assistant"
                and group.loc[i + 1, "Type"] == "code"
                and group.loc[i + 2, "Role"] == "assistant"
                and group.loc[i + 2, "Type"] == "websearch"
                and group.loc[i + 3, "Role"] == "assistant"
                and group.loc[i + 3, "Type"] == "websearch_result"
                and group.loc[i + 4, "Role"] == "assistant"
                and group.loc[i + 4, "Type"] == "text"
            ):
                websearch_results.append(
                    {
                        "问题": group.loc[i, "Content"],
                        "回答": group.loc[i + 4, "Content"],
                        "分发": "websearch",
                        "判定": "",
                    }
                )
                i += 5
            else:
                i += 1

    results = chat_results + websearch_results
    if results:
        df_export = pd.DataFrame(results)
        df_export.to_excel(OUTPUT_EXCEL, index=False)
        print(
            f"✅ 导出 {len(results)} 条记录（chat: {len(chat_results)}, websearch: {len(websearch_results)}）至 `{OUTPUT_EXCEL}`"
        )
    else:
        print("⚠️ 未匹配到符合结构的记录，请检查数据结构")


def read_excel_rows(file_path, sheet_name):
    try:
        wb = openpyxl.load_workbook(file_path)
        sheet = wb[sheet_name]
        headers = [cell.value for cell in sheet[1]]

        q_idx = headers.index("问题")
        a_idx = headers.index("回答")

        data = []
        for row_index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            data.append(
                {
                    "row_index": row_index,
                    "问题": row[q_idx].value or "",
                    "回答": row[a_idx].value or "",
                }
            )

        print(f"[INFO] 读取 Excel 成功，共 {len(data)} 条记录")
        return data

    except Exception as e:
        print(f"[ERROR] 读取 Excel 失败：{e}")
        return []


def write_cell(file_path, sheet_name, cell_position, content):
    try:
        wb = openpyxl.load_workbook(file_path)
        sheet = wb[sheet_name]
        sheet[cell_position] = content
        wb.save(file_path)
    except Exception as e:
        print(f"[ERROR] 写入单元格 {cell_position} 失败：{e}")


def run_one(row):
    try:
        prompt = Prompt.replace("{{.question}}", str(row.get("问题", ""))).replace(
            "{{.answer}}", str(row.get("回答", ""))
        )
        messages = [{"role": "user", "name": "user", "content": prompt}]
        api = GateWayAPI(retry_count=1)

        success, content, _ = api.chat_msgs(
            model, messages, context, llm_arguments, version, sec_text
        )
        if not success:
            raise RuntimeError("模型调用失败")

        cell = f"{输出列编号}{row['row_index']}"
        write_cell(OUTPUT_EXCEL, sheetname, cell, content.strip())
        return True
    except Exception as e:
        print(f"[ERROR] 第 {row['row_index']} 行分析失败：{e}")
        return False


def run_all():
    rows = read_excel_rows(OUTPUT_EXCEL, sheetname)
    total = len(rows)
    success_count, fail_count = 0, 0

    for idx, row in enumerate(rows, start=1):
        print(f"--- 分析第 {idx}/{total} 条 ---")
        if run_one(row):
            success_count += 1
        else:
            fail_count += 1

    print(f"\n✅ 分析完成：总共 {total} 条，成功 {success_count}，失败 {fail_count}")
    write_to_sqlite()


def write_to_sqlite():
    try:
        wb = openpyxl.load_workbook(OUTPUT_EXCEL)
        sheet = wb[sheetname]
        headers = [cell.value for cell in sheet[1]]

        q_idx = headers.index("问题")
        a_idx = headers.index("回答")
        f_idx = headers.index("分发") if "分发" in headers else None
        r_idx = headers.index("判定") if "判定" in headers else None

        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute("DROP TABLE IF EXISTS panding")
        cursor.execute(
            """
            CREATE TABLE panding (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                question TEXT,
                answer TEXT,
                dispatch TEXT,
                result TEXT
            )
        """
        )

        for row in sheet.iter_rows(min_row=2):
            cursor.execute(
                "INSERT INTO panding (question, answer, dispatch, result) VALUES (?, ?, ?, ?)",
                (
                    str(row[q_idx].value or ""),
                    str(row[a_idx].value or ""),
                    str(row[f_idx].value) if f_idx is not None else "",
                    str(row[r_idx].value) if r_idx is not None else "",
                ),
            )

        conn.commit()
        conn.close()
        print("[INFO] 已将评分结果写入 SQLite 数据库（表：panding）")
    except Exception as e:
        print(f"[ERROR] 写入 SQLite 数据库失败：{e}")


def analyze_score_distribution(df):
    if df.empty:
        return pd.DataFrame()

    counts = df["result"].value_counts().sort_index()
    summary = pd.DataFrame({"score": counts.index, "count": counts.values}).reset_index(
        drop=True
    )
    return summary


# 数据库路径和评分表名
DB_PATH = "dialog_data.db"
PANDING_TABLE = "panding"

# 设置中文字体及负号显示
matplotlib.rcParams["font.family"] = [
    "SimHei",
    "Microsoft YaHei",
    "Arial Unicode MS",
    "sans-serif",
]
matplotlib.rcParams["axes.unicode_minus"] = False


def check_panding_table_exist():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name='panding'"
    )
    result = cursor.fetchone()
    conn.close()
    return result is not None


def load_panding_data():
    conn = sqlite3.connect(DB_PATH)
    df = pd.read_sql_query("SELECT * FROM panding", conn)
    conn.close()
    return df


def plot_bar_and_avg(data, title):

    order = ["3", "2", "1", "0"]
    counts = data["result"].value_counts().reindex(order, fill_value=0)

    numeric_scores = pd.to_numeric(data["result"], errors="coerce")
    avg_score = numeric_scores.dropna().mean()

    fig, ax = plt.subplots(figsize=(5, 4))
    values = counts.values
    bars = ax.bar(
        counts.index, values, color=["#4CAF50", "#2196F3", "#FFC107", "#F44336"]
    )

    ax.set_title(title)
    ax.set_xlabel("评分")
    ax.set_ylabel("数量")
    ax.set_ylim(
        0, max(values) * 1.2 if max(values) > 0 else 1
    )  # 留白顶部，避免数字被遮挡
    ax.grid(axis="y", linestyle="--", alpha=0.5)

    for i, v in enumerate(values):
        ax.text(
            i,
            v + max(values) * 0.03,
            str(v),
            ha="center",
            va="bottom",
            fontsize=10,
            color="black",
        )

    # 可选：显示平均分水平线及标签
    if not pd.isna(avg_score):
        ax.axhline(avg_score, color="red", linestyle="--", linewidth=1)
        ax.text(
            len(counts) - 0.5,
            avg_score + max(values) * 0.02,
            f"平均分: {avg_score:.2f}",
            color="red",
            ha="right",
            va="bottom",
            fontsize=9,
        )

    fig.tight_layout()
    return fig, avg_score


def analyze_and_prepare_figs(df):
    import matplotlib.pyplot as plt
    import pandas as pd

    df["result"] = pd.to_numeric(df["result"], errors="coerce")
    avg_scores = df.groupby("dispatch")["result"].mean().to_dict()

    figs = {}
    for dispatch_type in df["dispatch"].unique():
        dispatch_df = df[df["dispatch"] == dispatch_type]

        counts = dispatch_df["result"].value_counts().sort_index()
        fig, ax = plt.subplots(figsize=(5, 4))

        bars = ax.bar(counts.index.astype(str), counts.values, color="#4CAF50")
        ax.set_title(f"{dispatch_type} 评分分布")
        ax.set_xlabel("分数")
        ax.set_ylabel("数量")
        ax.set_ylim(0, max(counts.values) * 1.2 if max(counts.values) > 0 else 1)
        ax.grid(axis="y", linestyle="--", alpha=0.5)

        for i, v in enumerate(counts.values):
            ax.text(
                i,
                v + max(counts.values) * 0.03,
                str(v),
                ha="center",
                va="bottom",
                fontsize=10,
            )

        fig.tight_layout()
        figs[dispatch_type] = fig

    return figs, avg_scores
